# -*- coding: utf-8 -*-
# BODA VMS·VMS.Web GS 인증 및 향후 일정 보고서 — FMS 세부개발일정 보고서와 동일 형식(5시트)
# v2: 로드맵 간트 조건부서식 수정 (DXF 는 start/end color 모두 필요) + 열 때 전체 재계산
import datetime as dt
import sys
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import CellIsRule
from openpyxl.utils import get_column_letter

OUT = sys.argv[1] if len(sys.argv) > 1 else r"C:\Users\vinos\Downloads\BODA_VMS_GS인증_세부일정_보고서.xlsx"

# ---- 팔레트 (FMS 문서와 동일 계열) ----
NAVY = "1F3864"; BLUE = "2E5496"
FILL_LABEL = "D9E1F2"; FILL_DONE = "EEF3FB"; FILL_TERM = "F2ECF7"; FILL_NOTE = "FFF2CC"
GRP = ["E2E3E5", "D8EFEA", "DEEAF6", "FCEFD6", "E1D5E7"]  # 그룹 행 색
BAR = "2E75B6"  # 간트 막대

thin = Side(style="thin", color="BFBFBF")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)

def F(color=None, bold=False, size=11, white=False):
    return Font(name="맑은 고딕", size=size, bold=bold,
                color="FFFFFF" if white else (color or "1A1A1A"))

def fill(hexv): return PatternFill("solid", fgColor=hexv)

WRAP = Alignment(vertical="center", wrap_text=True)
CTR = Alignment(horizontal="center", vertical="center", wrap_text=True)

def put(ws, coord, value, font=None, fl=None, align=None, border=True, numfmt=None):
    c = ws[coord]
    c.value = value
    c.font = font or F()
    if fl: c.fill = fill(fl)
    c.alignment = align or WRAP
    if border: c.border = BORDER
    if numfmt: c.number_format = numfmt
    return c

wb = openpyxl.Workbook()
wb.calculation.fullCalcOnLoad = True  # 열 때 전체 재계산 — 수식 캐시 없음 대응

# ============================================================ 1) 표지
ws = wb.active; ws.title = "표지"
ws.sheet_view.showGridLines = False
for col, w in {"A": 2.5, "B": 22, "C": 24, "D": 13, "E": 13, "F": 18}.items():
    ws.column_dimensions[col].width = w

ws.merge_cells("B2:F3")
put(ws, "B2", "BODA VMS · VMS.Web GS 인증 및 향후 일정 보고서", F(bold=True, size=18, white=True), NAVY, CTR)
ws.merge_cells("B4:F4")
put(ws, "B4", "GS(Good Software) 인증 신청 · 시험 대응 · 출시 준비 · 후속 개선 계획", F(bold=False, size=11, white=True), BLUE, CTR)

def section(row, title):
    ws.merge_cells(f"B{row}:F{row}")
    put(ws, f"B{row}", title, F(bold=True, size=12, white=True), BLUE, WRAP)

def kv(row, k, v, key_fill=FILL_LABEL):
    put(ws, f"B{row}", k, F(bold=True, size=10), key_fill)
    ws.merge_cells(f"C{row}:F{row}")
    put(ws, f"C{row}", v, F(size=10))

section(6, "1. 문서 정보")
kv(7,  "문서명", "BODA VMS · VMS.Web GS 인증 및 향후 일정 보고서")
kv(8,  "프로젝트명", "BODA VMS(현장 비전검사) + BODA.VMS.Web(관리 서버) GS 인증 획득")
kv(9,  "작성일", "2026-07-07")
kv(10, "작성자 / 부서", "(작성자)  /  R&D 사업부")
kv(11, "대상 기간", "2026-07-27 ~ 2026-12-25 · 약 22주 (자체 검증→신청→시험→인증 + 출시 준비·후속)")
kv(12, "문서 버전", "v1.0 (비전문가용 · 쉬운 용어)")

section(14, "2. 이번 보고의 내용 (한눈에)")
kv(15, "① GS 인증이란", "소프트웨어 품질을 국가 공인 시험기관(TTA)이 시험해 주는 국가 인증 — 1등급 목표")
kv(16, "② 준비 완료", "두 프로그램의 보안·품질 보강과 시험 제출 문서(매뉴얼 등 5종) 준비가 끝남")
kv(17, "③ 자체 검증→시험", "2주 자체 검증(시험 리허설) 후 신청서 제출 → 시험 대응 → 보완 → 인증 획득")
kv(18, "④ 출시 전 보완", "설치 파일 서명(보안 경고 제거), 납품 문안 정비(라이선스 초안 정리) 등 정식 배포 준비")
kv(19, "⑤ 후속 개선(선택)", "인증 이후 3D 카메라 연동·보안 로그 외부 전송·운영 편의 기능 고도화")

section(21, "3. 시스템 현황 (이미 완료되어 있음 — 이번 일정의 근거)")
kv(22, "VMS (현장 검사)", "보안·품질 보강 58건 완료 · 자동 테스트 571건+ 통과 · 감사기록/백업/보존 체계 구축", FILL_DONE)
kv(23, "VMS.Web (관리 서버)", "보강 36건 완료 · 자동 테스트 424건 통과 · 계정 승인/알람/이력 관리", FILL_DONE)
kv(24, "제출 문서 5종", "사용자매뉴얼 v1.1(89쪽·쉬운말 전면 개편) · 제품설명서 · OSS 확인서 · 신청서 · 체크리스트", FILL_DONE)
kv(25, "사용 편의", "4개 프로그램 전 화면에 ⓘ 도움말 적용 — 심사관·사용자가 화면에서 바로 설명 확인", FILL_DONE)

section(27, "4. 쉬운 용어 안내")
kv(28, "GS 인증", "정부가 지정한 시험기관이 소프트웨어 품질을 시험·인증하는 제도 (공공 납품에 유리)", FILL_TERM)
kv(29, "VMS", "현장 PC에서 카메라로 제품을 자동 검사하는 프로그램", FILL_TERM)
kv(30, "BODA.VMS.Web", "검사 결과·작업지시·알람을 한곳에서 관리하는 서버(웹) 프로그램", FILL_TERM)
kv(31, "코드 서명", "설치 파일에 제조사 인증 서명을 넣어 위변조와 보안 경고를 막는 것", FILL_TERM)
kv(32, "OSS 고지", "제품에 포함된 오픈소스 사용 내역을 고객에게 알리는 문서(NOTICE) — 납품 시 첨부", FILL_TERM)
kv(33, "SIEM 연동", "여러 시스템의 보안 기록을 한곳에 모아 감시하는 체계로 보내는 것", FILL_TERM)

section(35, "5. 참고 (일정 전제)")
kv(36, "일정 전제", "시험 기간은 TTA 접수 상황에 따라 변동(통상 4~8주). 날짜는 조정 가능한 추정치", FILL_NOTE)
kv(37, "결정 필요①", "GS 신청 제출 시점 확정 (시험 비용 결재 필요)", FILL_NOTE)
kv(38, "결정 필요②", "코드 서명 인증서 구매 (외부 첫 배포 전 필수 — 발급 1~2주 소요)", FILL_NOTE)
kv(39, "사용권 조건", "별도 사용권 계약(EULA) 없이 개별 납품계약에서 규정 — 라이선스 문안 표준화만 진행", FILL_NOTE)

ws.merge_cells("B41:F41")
put(ws, "B41", "※ 왜 하는지 → [배경·목적] 시트 / 상세 일정 → [세부일정] 시트 / 기간별 막대그림 → [로드맵] 시트 / 주요 목표 → [마일스톤] 시트",
    F(color="595959", size=10), None, WRAP, border=False)
ws.row_dimensions[41].height = 33.5

# ============================================================ 2) 배경·목적
ws = wb.create_sheet("배경·목적")
ws.sheet_view.showGridLines = False
for col, w in {"A": 2.5, "B": 21, "C": 39, "D": 37, "E": 35}.items():
    ws.column_dimensions[col].width = w
ws.merge_cells("B2:E2")
put(ws, "B2", "이 일정이 왜 필요한가 (배경·목적)", F(bold=True, size=14, white=True), NAVY, CTR)
ws.row_dimensions[2].height = 21
ws.merge_cells("B3:E3")
put(ws, "B3", "각 항목의 '현재 상황 → 목적 → 기대 효과'를 쉽게 정리했습니다.", F(color="595959"), None, WRAP, border=False)
for i, h in enumerate(["항목", "현재 상황·문제점", "목적 (왜 하는가)", "기대 효과"]):
    put(ws, f"{get_column_letter(2+i)}4", h, F(bold=True, white=True), BLUE, CTR)

bg_rows = [
    ("① GS 인증 획득   ★핵심", "E7E6F7",
     "공공기관·대기업 납품에 GS 인증을 요구하는 경우가 많은데, 제품 품질을 객관적으로 증명할 자료가 없음",
     "국가 공인 시험기관(TTA)의 시험을 통과해 품질을 공식 인증받음",
     "공공조달 등록 가능 · 영업 경쟁력 강화 · 제품 신뢰도 상승"),
    ("② 시험 대응 체계", "D8EFEA",
     "심사관은 매뉴얼을 보고 실제 동작을 확인함 — 문서와 제품이 다르면 지적(결함)으로 이어짐",
     "매뉴얼·제품설명서와 실제 화면/동작을 일치시키고, 지적 사항에 빠르게 대응할 준비",
     "시험 기간 단축 · 재시험 횟수 최소화"),
    ("③ 설치 파일 서명", "DEEAF6",
     "서명 없는 설치 파일은 고객 PC에서 '알 수 없는 게시자' 보안 경고가 떠 설치를 꺼리게 됨",
     "제조사 인증서로 설치 파일에 서명해 경고 없이 설치되게 함",
     "고객 설치 경험 개선 · 위변조 방지 · 기업 신뢰도"),
    ("④ 납품 문안 정비", "FCEFD6",
     "제품과 함께 나가는 라이선스 문서에 초안(DRAFT) 표기가 남아 있음 — 납품 문서로서 모양이 좋지 않음",
     "표준 문안으로 확정하고 오픈소스 고지(NOTICE) 첨부를 점검 — 사용권 조건 자체는 개별 납품계약에서 규정",
     "납품 문서 품질 확보 · 오픈소스 고지 의무 준수"),
    ("⑤ 후속 개선(선택)", "E1D5E7",
     "3D 카메라 일부 설정이 아직 실제 장비와 연결되지 않았고, 보안 기록의 외부 전송이 꺼져 있음",
     "판매 구성에 따라 3D 연동·보안 로그 외부 전송·운영 편의 기능을 순차 고도화",
     "제품 차별화 · 다중 사이트 운영 대응"),
]
r = 5
for title, c, cur, why, eff in bg_rows:
    put(ws, f"B{r}", title, F(bold=True, size=10), c)
    put(ws, f"C{r}", cur, F(size=10)); put(ws, f"D{r}", why, F(size=10)); put(ws, f"E{r}", eff, F(size=10))
    ws.row_dimensions[r].height = 52
    r += 1
ws.merge_cells(f"B{r+1}:E{r+1}")
put(ws, f"B{r+1}", "※ 요약: 개발·문서 준비는 끝났고, 이제 '신청 → 시험 대응 → 인증 → 정식 배포' 절차를 밟는 단계입니다.",
    F(bold=True, color="595959"), None, WRAP, border=False)

# ============================================================ 3) 세부일정 (WBS)
ws = wb.create_sheet("세부일정")
widths = {"A": 4, "B": 9, "C": 11, "D": 30, "E": 42, "F": 12, "G": 11, "H": 11, "I": 7, "J": 8, "K": 8, "L": 15, "M": 13}
for col, w in widths.items(): ws.column_dimensions[col].width = w
ws.freeze_panes = "D5"
ws.merge_cells("A1:M1")
put(ws, "A1", "BODA VMS · VMS.Web GS 인증 세부일정 (작업 분류 체계)", F(bold=True, size=14, white=True), NAVY, CTR)
ws.row_dimensions[1].height = 21
ws.merge_cells("A2:M2")
put(ws, "A2", "기간 2026-07-27 ~ 2026-12-25 (약 22주) · 작성일 2026-07-07 · 완료 항목은 근거로 함께 표기",
    F(color="595959", size=10), None, WRAP, border=False)
for i, h in enumerate(["No", "구분", "분야", "작업 항목(무엇을)", "상세 설명(어떻게)", "담당",
                       "시작일", "종료일", "기간(일)", "진행률", "상태", "결과물", "비고"]):
    put(ws, f"{get_column_letter(1+i)}4", h, F(bold=True, white=True), NAVY, CTR)
ws.row_dimensions[4].height = 34

D = dt.datetime
groups = [
    ("1. GS 인증 준비 (완료 — 근거)", GRP[0], [
        ("VMS", "보안·품질 보강 58건", "인증 기준에 맞춘 보안 8종·감사기록·백업/복원·보존정책 등 보강", "개발",
         D(2026, 3, 2), D(2026, 6, 4), 1.0, "완료", "보강 완료 (테스트 571+)", ""),
        ("Web", "관리 서버 보강 36건", "계정 승인·권한·알람·이력 등 인증 기준 보강", "개발",
         D(2026, 3, 2), D(2026, 6, 4), 1.0, "완료", "보강 완료 (테스트 424)", ""),
        ("문서", "GS 제출 문서 5종 작성", "사용자매뉴얼·제품설명서·OSS 확인서·신청서·체크리스트", "개발",
         D(2026, 5, 18), D(2026, 6, 12), 1.0, "완료", "제출 문서 5종", ""),
        ("문서", "사용자매뉴얼 쉬운말 전면 개편", "전 장을 현장 눈높이로 재작성 + 딥러닝 장 신설 (v1.1 · 89쪽) + 검토 반영", "개발",
         D(2026, 6, 30), D(2026, 7, 14), 0.8, "진행중", "매뉴얼 v1.1", "검토 반영 1주 연장"),
        ("공통", "전 화면 ⓘ 도움말 적용", "4개 프로그램 설정 항목·화면 옆에 쉬운 설명 말풍선 추가", "개발",
         D(2026, 7, 6), D(2026, 7, 7), 1.0, "완료", "도움말 적용판", ""),
    ]),
    ("2. 자체 검증 테스트 (GS 시험 리허설)", GRP[1], [
        ("계획", "자체 검증 계획·체크리스트 확정", "GS 시험 관점(기능적합성·사용성·신뢰성·보안성 등)으로 점검 항목을 정리", "개발·PM",
         D(2026, 7, 13), D(2026, 7, 14), 0, "예정", "검증 체크리스트", ""),
        ("VMS", "VMS 기능·매뉴얼 일치 검증", "매뉴얼 전 장을 그대로 따라 하며 실제 동작·화면·문구 일치 확인 (설치→설정→검사→관리)", "개발",
         D(2026, 7, 14), D(2026, 7, 17), 0, "예정", "검증 결과서(VMS)", "결함 즉시 기록"),
        ("Web", "VMS.Web 기능·연동 검증", "웹 전 화면 기능 + VMS 연동(결과 업로드·알람·파라미터 동기화) 확인", "개발",
         D(2026, 7, 15), D(2026, 7, 20), 0, "예정", "검증 결과서(Web)", ""),
        ("품질", "설치/제거·운영 회귀 검증", "설치 파일 설치·제거·업그레이드 + 운영 회귀 체크리스트(26항목) 통과", "개발",
         D(2026, 7, 20), D(2026, 7, 22), 0, "예정", "회귀 통과 기록", ""),
        ("보완", "발견 결함 수정·재확인", "자체 검증에서 나온 문제를 수정·재확인하고 제출 문서에 반영", "개발",
         D(2026, 7, 22), D(2026, 7, 24), 0, "예정", "수정 반영판", ""),
    ]),
    ("3. GS 신청 · 시험 대응", GRP[2], [
        ("신청", "신청 서류 최종 점검·제출", "제출 문서 5종 최종본 확정 후 TTA 에 신청 접수", "개발·PM",
         D(2026, 7, 13), D(2026, 7, 17), 0, "예정", "신청 접수증", "결정 필요①"),
        ("신청", "시험 계약·일정 확정", "시험기관과 계약, 시험 착수일 확정", "PM",
         D(2026, 7, 20), D(2026, 7, 31), 0, "예정", "시험 일정표", "TTA 상황 따라 변동"),
        ("환경", "시험 환경·설치본 준비", "시험용 PC·카메라 구성, 설치 파일 설치/제거 검증, 운영 회귀 26항목 사전 점검", "개발",
         D(2026, 7, 20), D(2026, 7, 31), 0, "예정", "시험 환경", ""),
        ("시험", "1차 시험 대응", "심사관 질의 응답·현상 재현 지원·기능 시연", "개발",
         D(2026, 8, 3), D(2026, 8, 28), 0, "예정", "시험 결과서", "기간은 접수 상황 따라"),
        ("보완", "지적사항(결함) 보완", "시험에서 나온 지적사항 수정 + 수정판·문서 갱신 제출", "개발",
         D(2026, 8, 31), D(2026, 9, 11), 0, "예정", "보완 수정판", ""),
        ("시험", "재확인 시험 대응", "보완 항목 재확인 시험 지원", "개발",
         D(2026, 9, 14), D(2026, 9, 25), 0, "예정", "재확인 결과", ""),
        ("인증", "인증 획득·등록", "인증서 수령, 조달청 등록 등 후속 행정", "PM",
         D(2026, 9, 28), D(2026, 10, 2), 0, "예정", "GS 인증서", ""),
    ]),
    ("4. 정식 배포(출시) 준비 — 시험과 병행", GRP[3], [
        ("보안", "코드 서명 인증서 구매·적용", "제조사 인증서 발급 후 설치 파일(MSI) 서명 — 보안 경고 제거", "개발",
         D(2026, 8, 3), D(2026, 8, 21), 0, "예정", "서명된 설치 파일", "결정 필요②"),
        ("보안", "빌드 자동 서명 구성(선택)", "빌드 서버에서 자동으로 서명된 설치 파일 산출", "개발",
         D(2026, 8, 24), D(2026, 9, 4), 0, "예정", "자동 서명 빌드", "선택"),
        ("문서", "납품 표준 문안 정리", "라이선스 문서의 초안(DRAFT) 표기 제거·표준 문안 확정 + 오픈소스 고지(NOTICE·확인서) 첨부 점검", "개발·PM",
         D(2026, 7, 20), D(2026, 7, 24), 0, "예정", "표준 문안", "사용권은 납품계약에서 규정"),
        ("품질", "실 운영 회귀 검증", "실제 운영 시나리오 체크리스트(26항목) 전체 통과 확인", "개발",
         D(2026, 9, 21), D(2026, 9, 25), 0, "예정", "회귀 통과 보고", ""),
        ("배포", "정식 배포판 v1.2 릴리스", "서명된 설치 파일 + 릴리스 노트 + 자동 업데이트 알림 배포", "개발",
         D(2026, 10, 5), D(2026, 10, 16), 0, "예정", "v1.2 배포판", "인증 획득 후"),
    ]),
    ("5. 후속 개선 (선택 — 인증 이후)", GRP[4], [
        ("보안", "보안 기록 외부 전송(SIEM) 활성", "여러 사이트의 보안 기록을 한곳에서 감시하도록 연동", "개발",
         D(2026, 10, 19), D(2026, 10, 30), 0, "예정", "SIEM 연동", "다중 사이트 고객용"),
        ("기능", "3D 카메라 설정 실연결", "필터 강도·측정 범위 등 3D 설정 값을 실제 장비 동작에 연결", "개발",
         D(2026, 10, 19), D(2026, 11, 13), 0, "예정", "3D 연동판", "3D 구성 판매 시"),
        ("기능", "운영 편의 기능 고도화", "검사 이력 보존 옵션·감사 화면 개선·보존정책 변경 이력 등", "개발",
         D(2026, 11, 2), D(2026, 11, 27), 0, "예정", "개선판", ""),
        ("설치", "설치 마법사에 보안·보존 설정 통합", "첫 설치 때 보안/백업/보존 설정을 한 번에 입력하도록 개선", "개발",
         D(2026, 11, 16), D(2026, 11, 27), 0, "예정", "마법사 개선", ""),
    ]),
]

# 순연 규칙 — 완료(1.) 0주 / 자체 검증(2.) 2주 / 이후 그룹(3.~5.) 4주
# = 기존 1주 순연 + 자체 검증 2주 삽입 + 매뉴얼 개편 연장 1주
SHIFT_DAYS = {"1.": 0, "2.": 14}
def _shift(title): return dt.timedelta(days=SHIFT_DAYS.get(title[:2], 28))
groups = [
    (gtitle, gcolor,
     [(cat, what, how, owner, s + _shift(gtitle), e + _shift(gtitle), prog, status, outp, note)
      for (cat, what, how, owner, s, e, prog, status, outp, note) in items])
    for (gtitle, gcolor, items) in groups
]

row = 5
no = 1
wbs_positions = []
for gtitle, gcolor, items in groups:
    grow = row
    first, last = row + 1, row + len(items)
    ws.merge_cells(f"B{grow}:F{grow}")
    put(ws, f"B{grow}", gtitle, F(bold=True), gcolor)
    put(ws, f"A{grow}", None, fl=gcolor)
    put(ws, f"G{grow}", f"=MIN(G{first}:G{last})", F(bold=True), gcolor, CTR, numfmt="yyyy-mm-dd")
    put(ws, f"H{grow}", f"=MAX(H{first}:H{last})", F(bold=True), gcolor, CTR, numfmt="yyyy-mm-dd")
    put(ws, f"I{grow}", f"=H{grow}-G{grow}+1", F(bold=True), gcolor, CTR)
    put(ws, f"J{grow}", f"=IFERROR(AVERAGE(J{first}:J{last}),0)", F(bold=True), gcolor, CTR, numfmt="0%")
    for col in "KLM": put(ws, f"{col}{grow}", None, fl=gcolor)
    row += 1
    item_rows = []
    for (cat, what, how, owner, s, e, prog, status, outp, note) in items:
        put(ws, f"A{row}", no, align=CTR); put(ws, f"C{row}", cat, F(size=10), align=CTR)
        put(ws, f"D{row}", what, F(size=10)); put(ws, f"E{row}", how, F(size=10))
        put(ws, f"F{row}", owner, F(size=10), align=CTR)
        put(ws, f"G{row}", s, align=CTR, numfmt="yyyy-mm-dd")
        put(ws, f"H{row}", e, align=CTR, numfmt="yyyy-mm-dd")
        put(ws, f"I{row}", f"=H{row}-G{row}+1", align=CTR)
        put(ws, f"J{row}", prog, align=CTR, numfmt="0%")
        st_font = F(color="2E7D32", bold=True, size=10) if status == "완료" else (F(color="1565C0", size=10) if status == "진행중" else F(size=10))
        put(ws, f"K{row}", status, st_font, align=CTR)
        put(ws, f"L{row}", outp, F(size=10)); put(ws, f"M{row}", note, F(color="C00000" if "결정" in note else "595959", size=10))
        put(ws, f"B{row}", None)
        ws.row_dimensions[row].height = 32
        item_rows.append((no, what, owner, s, e, status))
        no += 1; row += 1
    wbs_positions.append((gtitle, gcolor, item_rows))

# ============================================================ 4) 로드맵 (주 단위 간트)
ws = wb.create_sheet("로드맵")
ws.sheet_view.showGridLines = False
for col, w in {"A": 4, "B": 9, "C": 32, "D": 12, "E": 11, "F": 11, "G": 6}.items():
    ws.column_dimensions[col].width = w

weeks = []
d = dt.date(2026, 7, 6)  # 완료 그룹(6월 말~7월 초) 꼬리도 보이도록 앞에서 시작
while d <= dt.date(2026, 12, 25):
    weeks.append(d); d += dt.timedelta(days=7)
first_week_col = 8  # H
for i in range(len(weeks)):
    ws.column_dimensions[get_column_letter(first_week_col + i)].width = 4.75

ws.merge_cells(f"A1:{get_column_letter(7 + len(weeks))}1")
put(ws, "A1", "BODA VMS · VMS.Web GS 인증 로드맵 (기간별 막대그림)", F(bold=True, size=14, white=True), NAVY, CTR)
ws.row_dimensions[1].height = 21
ws.merge_cells(f"A2:{get_column_letter(7 + len(weeks))}2")
put(ws, "A2", "주(월요일) 단위 · 파란 막대 = 작업 기간 · 시작/종료일 변경 시 막대 자동 갱신",
    F(color="595959", size=10), None, WRAP, border=False)

for i, h in enumerate(["No", "구분", "작업 항목", "담당", "시작일", "종료일", "기간"]):
    col = get_column_letter(1 + i)
    ws.merge_cells(f"{col}4:{col}5")
    put(ws, f"{col}4", h, F(bold=True, white=True), NAVY, CTR)
month_start = 0
for i in range(len(weeks) + 1):
    if i == len(weeks) or (i > 0 and weeks[i].month != weeks[month_start].month):
        c1 = get_column_letter(first_week_col + month_start)
        c2 = get_column_letter(first_week_col + i - 1)
        if c1 != c2: ws.merge_cells(f"{c1}4:{c2}4")
        put(ws, f"{c1}4", f"2026-{weeks[month_start].month:02d}", F(bold=True, white=True), NAVY, CTR)
        month_start = i
for i, wk in enumerate(weeks):
    put(ws, f"{get_column_letter(first_week_col + i)}5", dt.datetime(wk.year, wk.month, wk.day),
        F(bold=True, size=8, white=True), BLUE, CTR, numfmt="mm-dd")

def gantt_formulas(r):
    for i in range(len(weeks)):
        cl = get_column_letter(first_week_col + i)
        put(ws, f"{cl}{r}", f'=IF(AND($E{r}<={cl}$5+6,$F{r}>={cl}$5),1,"")', F(size=8, color="FFFFFF"), align=CTR)

r = 6
for gtitle, gcolor, item_rows in wbs_positions:
    grow = r
    first, last = r + 1, r + len(item_rows)
    ws.merge_cells(f"B{grow}:D{grow}")
    put(ws, f"B{grow}", gtitle, F(bold=True), gcolor)
    put(ws, f"A{grow}", None, fl=gcolor)
    put(ws, f"E{grow}", f"=MIN(E{first}:E{last})", F(bold=True), gcolor, CTR, numfmt="yyyy-mm-dd")
    put(ws, f"F{grow}", f"=MAX(F{first}:F{last})", F(bold=True), gcolor, CTR, numfmt="yyyy-mm-dd")
    put(ws, f"G{grow}", f"=F{grow}-E{grow}+1", F(bold=True), gcolor, CTR)
    gantt_formulas(grow)
    r += 1
    for (no_, what, owner, s, e, status) in item_rows:
        put(ws, f"A{r}", no_, align=CTR)
        put(ws, f"C{r}", what, F(size=10))
        put(ws, f"D{r}", owner, F(size=10), align=CTR)
        put(ws, f"E{r}", s, align=CTR, numfmt="yyyy-mm-dd")
        put(ws, f"F{r}", e, align=CTR, numfmt="yyyy-mm-dd")
        put(ws, f"G{r}", f"=F{r}-E{r}+1", align=CTR)
        put(ws, f"B{r}", None)
        gantt_formulas(r)
        r += 1

# DXF(조건부서식) 채움은 start/end color 를 모두 지정해야 Excel 이 칠한다 (openpyxl 특성)
bar_fill = PatternFill(start_color=BAR, end_color=BAR, fill_type="solid")
gantt_range = f"{get_column_letter(first_week_col)}6:{get_column_letter(first_week_col + len(weeks) - 1)}{r - 1}"
ws.conditional_formatting.add(gantt_range,
    CellIsRule(operator="equal", formula=["1"], fill=bar_fill, font=Font(color=BAR)))
ws.freeze_panes = "E6"

# ============================================================ 5) 마일스톤
ws = wb.create_sheet("마일스톤")
ws.sheet_view.showGridLines = False
for col, w in {"A": 3, "B": 6, "C": 30, "D": 13, "E": 16, "F": 38, "G": 9}.items():
    ws.column_dimensions[col].width = w
ws.merge_cells("B2:G2")
put(ws, "B2", "BODA VMS · VMS.Web GS 인증 주요 목표(마일스톤)", F(bold=True, size=14, white=True), NAVY, CTR)
ws.row_dimensions[2].height = 21
for i, h in enumerate(["No", "목표", "목표일", "관련 단계", "완료 기준", "상태"]):
    put(ws, f"{get_column_letter(2 + i)}3", h, F(bold=True, white=True), BLUE, CTR)

# 날짜는 순연 반영된 최종값 (세부일정과 일치: 그룹2 +2주, 그룹3~5 +4주)
milestones = [
    ("인증 준비 완료 (보강·문서·매뉴얼)", D(2026, 7, 14), "1", "보강 94건·테스트 995건·제출 문서 5종·매뉴얼 v1.1 검토 반영", "진행중"),
    ("자체 검증 테스트 완료",              D(2026, 8, 7),  "2", "매뉴얼-동작 일치·연동·회귀 통과, 발견 결함 수정 반영", "예정"),
    ("GS 신청 제출",                      D(2026, 8, 14), "3", "신청 서류 최종본 제출·접수 확인", "예정"),
    ("시험 착수",                          D(2026, 8, 31), "3", "계약·일정 확정, 시험 환경 인계", "예정"),
    ("1차 시험 완료",                      D(2026, 9, 25), "3", "1차 시험 종료·지적사항 목록 수령", "예정"),
    ("지적사항 보완 제출",                 D(2026, 10, 9), "3", "수정판 + 갱신 문서 제출", "예정"),
    ("GS 인증 획득",                       D(2026, 10, 30),"3", "인증서 수령", "예정"),
    ("정식 배포판 v1.2 출시",              D(2026, 11, 13),"4", "서명된 설치 파일·납품 문안 확정·릴리스 노트", "예정"),
    ("후속 개선 완료(선택)",               D(2026, 12, 25),"5", "SIEM 연동·3D 실연결·운영 편의 개선", "예정"),
]
r = 4
for i, (goal, dd, stage, crit, status) in enumerate(milestones, 1):
    put(ws, f"B{r}", i, F(size=10), align=CTR)
    put(ws, f"C{r}", goal, F(size=10))
    put(ws, f"D{r}", dd, F(size=10), align=CTR, numfmt="yyyy-mm-dd")
    put(ws, f"E{r}", stage, F(size=10), align=CTR)
    put(ws, f"F{r}", crit, F(size=10))
    st_font = F(color="2E7D32", bold=True, size=10) if status == "완료" else (F(color="1565C0", size=10) if status == "진행중" else F(size=10))
    put(ws, f"G{r}", status, st_font, align=CTR)
    r += 1

wb.save(OUT)
print("WROTE", OUT)
